from pathlib import Path

import pytest

from app.services import ingestion
from app.services.ingestion import build_chunks, humanize, load_sections, slugify

ROOT = Path(__file__).resolve().parents[1]


def test_slugify_and_humanize():
    assert slugify("Company HR Handbook (v2).pdf") == "company-hr-handbook-v2"
    assert slugify("***.md") == "document"
    assert humanize("company_hr_handbook.md") == "Company HR Handbook"
    assert humanize("pto-faq.docx") == "PTO FAQ"


def test_markdown_is_split_by_headings(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("# Title\n\n## Annual Leave\n20 days.\n\n## Sick Leave\n10 days.\n", encoding="utf-8")
    sections = load_sections(path)
    assert [(s.title, s.text) for s in sections] == [("Annual Leave", "20 days."), ("Sick Leave", "10 days.")]


def test_chunks_carry_section_title_and_metadata(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("## Remote Work\n" + ("Employees may work remotely two days per week. " * 60), encoding="utf-8")
    texts, metas = build_chunks(load_sections(path), "policy.md")
    assert len(texts) > 1
    assert all(t.startswith("Remote Work\n") for t in texts)
    assert all(m == {"source": "policy.md", "section": "Remote Work"} for m in metas)


def test_txt_and_docx_loading(tmp_path):
    txt = tmp_path / "notes.txt"
    txt.write_text("Payroll is paid monthly.", encoding="utf-8")
    assert load_sections(txt)[0].text == "Payroll is paid monthly."

    from docx import Document

    docx_path = tmp_path / "benefits.docx"
    document = Document()
    document.add_heading("Benefits", level=1)
    document.add_paragraph("Medical plan for full-time staff.")
    document.save(docx_path)
    sections = load_sections(docx_path)
    assert sections[0].title == "Benefits"
    assert "Medical plan" in sections[0].text


def test_pdf_loading_keeps_page_numbers():
    pdf = ROOT / "uploads" / "novacore_company_information.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not available")
    sections = load_sections(pdf)
    assert sections and sections[0].page == 1


def test_unsupported_type_raises(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("a,b", encoding="utf-8")
    with pytest.raises(ValueError):
        load_sections(path)


def test_ingest_file_replaces_previous_version_and_registers(tmp_path, monkeypatch):
    calls = {}
    monkeypatch.setattr(ingestion.vectorstore, "delete_document", lambda doc_id, known=0: calls.setdefault("deleted", doc_id))

    def fake_upsert(doc_id, texts, metadatas):
        calls["upserted"] = (doc_id, len(texts), metadatas[0])
        return len(texts)

    monkeypatch.setattr(ingestion.vectorstore, "upsert_chunks", fake_upsert)
    path = tmp_path / "leave_policy.md"
    path.write_text("## Annual Leave\n20 days per year.\n", encoding="utf-8")
    record = ingestion.ingest_file(path, origin="upload")

    assert calls["deleted"] == "leave-policy"
    assert calls["upserted"][0] == "leave-policy"
    assert calls["upserted"][2]["title"] == "Leave Policy"
    assert record["chunks"] == 1 and record["file_type"] == "md"
    assert any(d["doc_id"] == "leave-policy" for d in ingestion.storage.list_documents())


def test_empty_document_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(ingestion.vectorstore, "delete_document", lambda doc_id, known=0: 0)
    path = tmp_path / "empty.txt"
    path.write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError):
        ingestion.ingest_file(path)


def test_starter_knowledge_base_is_present():
    folder = ROOT / "data" / "knowledge_base"
    files = [p for p in folder.iterdir() if p.suffix in ingestion.SUPPORTED]
    assert len(files) >= 2
    for path in files:
        texts, _ = build_chunks(load_sections(path), path.name)
        assert texts


def test_delete_removes_known_chunk_ids_even_when_listing_lags(monkeypatch):
    deleted = []

    class FakeIndex:
        def delete(self, ids, namespace):
            deleted.extend(ids)

    monkeypatch.setattr(ingestion.vectorstore, "list_ids", lambda prefix=None: [])
    monkeypatch.setattr(ingestion.vectorstore, "get_index", lambda: FakeIndex())
    assert ingestion.vectorstore.delete_document("fresh-doc", known_chunks=3) == 3
    assert deleted == ["fresh-doc#0000", "fresh-doc#0001", "fresh-doc#0002"]
